"""Deterministic end-to-end integration pipeline test."""

from pathlib import Path
import openpyxl

from src.excel_handler import load_meridian_workbook, get_real_vendors_and_student_columns, to_vendor_record, export_inventory
from src.scoring import score_vendors
from src.osint import discover_sources_for_vendors
from src.classify import classify_text
from src.risk import analyse_risk
from src.meridian_export import build_meridian_updates
from src import db

WORKBOOK_PATH = Path("supporting material/Meridian_Vendor_Input.xlsx")


def test_end_to_end_pipeline(tmp_path: Path) -> None:
    # 1. Load Meridian Workbook
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    real_records, student_cols = get_real_vendors_and_student_columns(bundle)
    assert len(real_records) == 6
    vendors = [to_vendor_record(r) for r in real_records]

    # 2. Score Criticality
    criticality = score_vendors(vendors)
    assert len(criticality) == 6
    crit_by_name = {c.vendor_name: c for c in criticality}
    # Fiserv has critical dependency
    assert crit_by_name["Fiserv, Inc."].level.lower() == "critical"

    # 3. Discover Sources
    sources = discover_sources_for_vendors(vendors)
    assert len(sources) > 0

    # 4. Controlled Evidence Fixtures (Simulating deterministic OSINT retrieval)
    # Vendor 1: AutomWorx - no AI evidence (unrated)
    # Vendor 2: Fiserv - AI feature disclosure
    evidence = [
        classify_text(
            "Fiserv deploys generative AI and machine learning assistant capabilities within financial workflows.",
            source_url="https://www.fiserv.com/ai",
            source_title="Fiserv AI",
            vendor_name="Fiserv, Inc.",
        )[0],
        classify_text(
            "We maintain SOC 2 compliance and AES-256 encryption.",
            source_url="https://www.automworx.com/security",
            source_title="AutomWorx Security",
            vendor_name="AutomWorx",
        )[0],
    ]

    # 5. Risk Assessment
    evidence_by_vendor = {"Fiserv, Inc.": [evidence[0]], "AutomWorx": [evidence[1]]}
    risks = [
        analyse_risk(
            vendor_name=v.name,
            criticality=crit_by_name.get(v.name),
            evidence=evidence_by_vendor.get(v.name, []),
        )
        for v in vendors
    ]
    assert len(risks) == 6

    # Verify conservative risk rules
    automworx_risk = next(r for r in risks if r.vendor_name == "AutomWorx")
    # AutomWorx only has security_control evidence -> must be not_rated / insufficient_evidence
    assert automworx_risk.risk_level == "not_rated"
    assert automworx_risk.status == "insufficient_evidence"

    fiserv_risk = next(r for r in risks if r.vendor_name == "Fiserv, Inc.")
    # Fiserv has ai_capability_claim -> gets draft rating
    assert fiserv_risk.risk_score is not None
    assert fiserv_risk.status in ("heuristic_draft", "confirmed_disclosure", "potential_signal")

    # 6. Database Persistence
    db_file = tmp_path / "test_pipeline.db"
    conn = db.connect(db_file)
    for v in vendors:
        v_id = db.upsert_vendor(conn, v)
        v_src = [s for s in sources if s.vendor_name == v.name]
        db.save_sources(conn, v_id, v_src)
        v_ev = evidence_by_vendor.get(v.name, [])
        db.save_evidence(conn, v_id, v_ev)
        crit_item = crit_by_name.get(v.name)
        risk_item = next(r for r in risks if r.vendor_name == v.name)
        db.save_assessment(conn, vendor_id=v_id, criticality=crit_item, risk=risk_item)
    
    assessments = db.list_assessments(conn)
    assert len(assessments) == 6
    conn.close()

    # 7. Excel Export
    export_file = tmp_path / "exported_assessment.xlsx"
    updates = build_meridian_updates(
        bundle,
        vendors=vendors,
        criticality=criticality,
        sources=sources,
        evidence=evidence,
        risks=risks,
    )
    export_inventory(bundle, export_file, updates=updates)

    # 8. Re-open and Verify Export Integrity
    wb_exported = openpyxl.load_workbook(export_file)
    ws = wb_exported["Vendor Inventory"]
    # Check that student columns are populated for V-001 (row 6)
    crit_tier_val = ws.cell(6, 12).value  # Column L is Criticality Tier
    assert crit_tier_val is not None
    assert "High" in str(crit_tier_val) or "Critical" in str(crit_tier_val) or "Medium" in str(crit_tier_val)

    # Check that V-000 example row (row 5) was untouched
    example_id = ws.cell(5, 2).value
    assert example_id == "V-000"
    wb_exported.close()
