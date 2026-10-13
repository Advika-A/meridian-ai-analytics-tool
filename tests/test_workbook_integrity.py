"""Mandatory workbook integrity tests as required by Phase 3 & Section 6."""

import openpyxl
from pathlib import Path
import pytest

from src.excel_handler import (
    load_meridian_workbook,
    get_real_vendors_and_student_columns,
    export_inventory,
    ExcelHandlerError,
)
from src.models import CriticalityResult, EvidenceItem, RiskResult, SourceCandidate, VendorRecord
from src.meridian_export import build_meridian_updates

WORKBOOK_PATH = Path("supporting material/Meridian_Vendor_Input.xlsx")


def test_worksheets_and_order() -> None:
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    expected_sheets = ["Vendor Inventory", "Field Guide", "Points to consider"]
    assert wb.sheetnames == expected_sheets
    wb.close()


def test_vendor_inventory_header_row_and_columns() -> None:
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    assert bundle.header_row == 4
    col_names = [col.name for col in bundle.columns]
    assert "Vendor ID" in col_names
    assert "Vendor Name" in col_names
    assert "Criticality Tier" in col_names
    assert "AI Usage Detected (Y/N)" in col_names


def test_real_vendors_and_example_exclusion() -> None:
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    real_records, student_cols = get_real_vendors_and_student_columns(bundle)
    # Exactly 6 real vendors
    assert len(real_records) == 6
    real_ids = [r.vendor_id for r in real_records]
    assert real_ids == ["V-001", "V-002", "V-003", "V-004", "V-005", "V-006"]
    # V-000 is the fictional example and excluded from real records
    assert "V-000" not in real_ids
    example_records = [r for r in bundle.records if r.is_example]
    assert len(example_records) == 1
    assert example_records[0].vendor_id == "V-000"


def test_fictional_example_content_preserved_on_export(tmp_path: Path) -> None:
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    out_file = tmp_path / "export_test.xlsx"
    
    # Read original row 5 (V-000) values
    wb_orig = openpyxl.load_workbook(WORKBOOK_PATH)
    ws_orig = wb_orig["Vendor Inventory"]
    orig_row5 = [ws_orig.cell(5, col).value for col in range(1, ws_orig.max_column + 1)]
    wb_orig.close()

    # Export with dummy updates
    updates = {"V-001": {"Criticality Tier": "High"}}
    export_inventory(bundle, out_file, updates=updates)

    # Reopen and compare row 5
    wb_new = openpyxl.load_workbook(out_file)
    ws_new = wb_new["Vendor Inventory"]
    new_row5 = [ws_new.cell(5, col).value for col in range(1, ws_new.max_column + 1)]
    wb_new.close()

    assert new_row5 == orig_row5


def test_protected_cells_cannot_be_overwritten(tmp_path: Path) -> None:
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    out_file = tmp_path / "protected_test.xlsx"
    
    # Attempting to overwrite a protected column like 'Vendor Name' or 'Category'
    updates = {"V-001": {"Vendor Name": "Hacked Name", "Criticality Tier": "High"}}
    export_inventory(bundle, out_file, updates=updates)

    wb = openpyxl.load_workbook(out_file)
    ws = wb["Vendor Inventory"]
    # Row 6 is V-001
    assert ws.cell(6, 3).value == "AutomWorx"  # Preserved original Vendor Name!
    wb.close()


def test_repeated_export_stability(tmp_path: Path) -> None:
    bundle = load_meridian_workbook(WORKBOOK_PATH)
    out1 = tmp_path / "run1.xlsx"
    out2 = tmp_path / "run2.xlsx"
    
    updates = {"V-001": {"Criticality Tier": "High"}}
    export_inventory(bundle, out1, updates=updates)
    
    bundle2 = load_meridian_workbook(out1)
    export_inventory(bundle2, out2, updates=updates)

    wb1 = openpyxl.load_workbook(out1)
    wb2 = openpyxl.load_workbook(out2)
    assert wb1.sheetnames == wb2.sheetnames
    assert wb1["Vendor Inventory"].max_row == wb2["Vendor Inventory"].max_row
    wb1.close()
    wb2.close()


def test_missing_or_invalid_input_error(tmp_path: Path) -> None:
    fake_file = tmp_path / "nonexistent.xlsx"
    with pytest.raises(ExcelHandlerError):
        load_meridian_workbook(fake_file)
